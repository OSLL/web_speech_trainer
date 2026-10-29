"""Browser checks against the real web/worker/broker/Whisper compose stack."""

from pathlib import Path
import time
from urllib.parse import urlsplit
from uuid import uuid4

import pytest
import requests
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium_session import ROOT_DIR

pytestmark = pytest.mark.selenium


def api(selenium_session):
    client = requests.Session()
    host = urlsplit(selenium_session.host).hostname
    # CookieJar treats a single-label Docker hostname as <host>.local.
    domain = host if "." in host or ":" in host else host + ".local"
    for cookie in selenium_session.driver.get_cookies():
        client.cookies.set(
            cookie["name"], cookie["value"],
            domain=domain, path=cookie.get("path", "/"),
        )
    return client


def test_invalid_id_returns_404_in_browser(selenium_session):
    result = selenium_session.driver.execute_async_script("""
        const done=arguments[arguments.length-1];
        fetch('/api/trainings/not-an-id/').then(r=>done(r.status));
    """)
    assert result == 404


def test_upload_error_allows_retry_and_valid_upload(selenium_session, tmp_path):
    driver = selenium_session.driver
    bad = tmp_path / "invalid.pdf"
    bad.write_text("not a PDF")
    driver.get(selenium_session.host + "/upload_presentation/")
    upload = driver.find_element(By.ID, "file-loading")
    upload.send_keys(str(bad))
    button = driver.find_element(By.ID, "button-submit")
    button.click()
    WebDriverWait(driver, 15).until(
        lambda _: driver.find_element(By.ID, "alert").is_displayed()
    )
    assert button.is_enabled()
    assert driver.find_element(By.ID, "error-text").text
    upload.send_keys(str(Path(ROOT_DIR) / "test_data/test_presentation_file_0.pdf"))
    button.click()
    WebDriverWait(driver, 30).until(lambda d: "/trainings/" in d.current_url)
    assert driver.find_element(By.ID, "record").is_displayed()


def test_real_pipeline_duplicate_start_and_admin_recheck(selenium_session):
    driver = selenium_session.driver
    driver.get(selenium_session.host + "/init/")
    client = api(selenium_session)
    with open(
        Path(ROOT_DIR) / "test_data/test_presentation_file_0.pdf", "rb"
    ) as source:
        uploaded = client.post(
            selenium_session.host + "/api/files/presentations/",
            files={"presentation": ("test.pdf", source, "application/pdf")},
        )
    assert uploaded.status_code == 200 and uploaded.json()["message"] == "OK"
    created = client.post(
        selenium_session.host
        + f"/api/trainings/presentations/{uploaded.json()['presentation_file_id']}/"
    )
    assert created.status_code == 200, created.text
    training = created.json()["training_id"]
    for timestamp in (time.time(), time.time() + 15):
        response = client.put(
            selenium_session.host + f"/api/trainings/timestamps/{training}/",
            params={"timestamp": timestamp},
        )
        assert response.status_code == 200, response.text
    with open(Path(ROOT_DIR) / "simple_phrases_russian.mp3", "rb") as source:
        response = client.post(
            selenium_session.host + f"/api/trainings/presentation-records/{training}/",
            files={"presentationRecord": ("record.mp3", source, "audio/mpeg")},
            data={"presentationRecordDuration": "15"},
        )
    assert response.status_code == 200, response.text
    base = selenium_session.host + f"/api/trainings/{training}/"
    assert client.post(base).status_code == 200
    first = client.get(base + "workflow/").json()["run_id"]
    assert client.post(base).status_code == 200
    assert client.get(base + "workflow/").json()["run_id"] == first

    def complete(_):
        state = client.get(base + "workflow/").json()
        assert state["state"] != "failed", state
        return state if state["state"] == "complete" else False

    final = WebDriverWait(driver, 240, poll_frequency=2).until(complete)
    assert len(final["completed_stages"]) == 6
    driver.get(selenium_session.host + f"/trainings/statistics/{training}/")
    assert "Оценка за тренировку" in driver.find_element(By.ID, "feedback").text
    driver.execute_script(
        "setCriteriaResults(arguments[0])", '<img src=x onerror="window.feedbackXss=1">'
    )
    assert not driver.find_elements(By.CSS_SELECTOR, "#criteria-results img")
    assert not driver.execute_script("return window.feedbackXss || false")
    assert client.post(base + "?restart=1").status_code == 200
    assert client.get(base + "workflow/").json()["run_id"] != first
    WebDriverWait(driver, 240, poll_frequency=2).until(complete)


def test_assigned_presentation_starts_only_after_button_click(selenium_session):
    from app.config import Config

    driver = selenium_session.driver
    client = api(selenium_session)
    with open(
        Path(ROOT_DIR) / "test_data/test_presentation_file_0.pdf", "rb"
    ) as source:
        response = client.post(
            selenium_session.host + "/api/files/presentations/",
            files={"presentation": ("assigned.pdf", source, "application/pdf")},
        )
    assert response.status_code == 200
    launch = client.post(
        selenium_session.host + "/lti",
        data={
            "oauth_consumer_key": Config.c.testing.oauth_consumer_key,
            "ext_user_username": Config.c.testing.session_id,
            "lis_person_name_full": "Assigned presentation student",
            "custom_task_id": f"assigned-browser-{uuid4().hex}",
            "lis_outcome_service_url": Config.c.testing.lis_outcome_service_url,
            "lis_result_sourcedid": Config.c.testing.lis_result_source_did,
            "custom_attempt_count": "3",
            "custom_criteria_pack_id": "SimplePack",
            "custom_presentation_id": response.json()["presentation_file_id"],
            "roles": "Instructor",
        },
    )
    assert launch.status_code == 200
    driver.get(selenium_session.host + "/")
    for cookie in client.cookies:
        driver.add_cookie({"name": cookie.name, "value": cookie.value, "path": "/"})
    driver.get(selenium_session.host + "/training_greeting/")
    driver.refresh()
    assert (
        api(selenium_session)
        .get(selenium_session.host + "/api/task-attempts/")
        .json()["training_number"]
        == 0
    )
    driver.find_element(By.ID, "start-training-button").click()
    WebDriverWait(driver, 15).until(lambda d: "/trainings/" in d.current_url)
    assert driver.find_element(By.ID, "record").is_displayed()
    assert (
        api(selenium_session)
        .get(selenium_session.host + "/api/task-attempts/")
        .json()["training_number"]
        == 1
    )


def test_interview_upload_record_and_results(selenium_session, tmp_path):
    from docx import Document
    from app.config import Config

    driver = selenium_session.driver
    host = selenium_session.host
    task = 'interview-' + uuid4().hex
    client = requests.Session()
    launch = client.post(host + '/lti', data={
        'lis_person_name_full': 'Interview Test', 'ext_user_username': 'interview-test',
        'oauth_consumer_key': Config.c.testing.oauth_consumer_key,
        'custom_task_id': task, 'custom_mode': 'interview',
        'custom_interview_question_count': '1', 'custom_interview_attempt_count': '2',
        'roles': 'Learner', 'lis_outcome_service_url': 'http://example.com/outcomes',
        'lis_result_sourcedid': task,
    }, timeout=30)
    assert launch.status_code == 200, launch.text
    driver.get(host + '/')
    for cookie in client.cookies:
        driver.add_cookie({'name': cookie.name, 'value': cookie.value, 'path': '/'})
    driver.get(host + '/interview/upload/')
    wait = WebDriverWait(driver, 180, poll_frequency=1)
    wait.until(lambda d: d.find_element(By.ID, 'upload-mode').is_displayed())
    document = Document()
    document.add_heading('Введение', level=1)
    document.add_paragraph('Объект исследования — система обучения. Практическая значимость работы состоит в обучении студентов.')
    document.add_heading('Метод решения', level=1)
    document.add_paragraph('Система анализирует речь и формирует обратную связь.')
    document.add_heading('Заключение', level=1)
    document.add_paragraph('Разработана система обучения студентов.')
    path = tmp_path / 'interview.docx'
    document.save(path)
    driver.find_element(By.ID, 'document').send_keys(str(path))
    driver.find_element(By.ID, 'upload-submit-btn').click()
    wait.until(lambda d: d.current_url.rstrip('/').endswith('/interview'))
    wait.until(lambda d: d.find_element(By.ID, 'main-btn').is_enabled())
    assert driver.find_element(By.ID, 'main-btn').text == 'Начать'
    driver.find_element(By.ID, 'main-btn').click()
    wait.until(lambda d: d.find_element(By.ID, 'next-btn').is_displayed())
    assert driver.find_element(By.ID, 'question-text').text
    # Let the browser capture a meaningful section of the fake speech audio.
    started = time.monotonic()
    WebDriverWait(driver, 15, poll_frequency=1).until(lambda _: time.monotonic() - started >= 8)
    driver.find_element(By.ID, 'next-btn').click()
    wait.until(lambda d: '/interview/results/' in d.current_url)
    wait.until(lambda d: len(d.find_elements(By.CSS_SELECTOR, '#results-body tr')) == 3)
    assert '/' in driver.find_element(By.ID, 'total-score').text
    assert driver.find_element(By.ID, 'results-verdict').text
    recording_id = driver.current_url.rstrip('/').rsplit('/', 1)[1]
    feedback = client.get(host + '/api/interview/feedback/' + recording_id, timeout=20)
    assert feedback.status_code == 200 and 0 <= feedback.json()['score'] <= 1
    # A different LMS namespace with the same username/task cannot read this result.
    foreign = requests.Session()
    result = foreign.post(host + '/lti', data={
        'lis_person_name_full': 'Other LMS', 'ext_user_username': 'interview-test',
        'oauth_consumer_key': 'other-' + uuid4().hex, 'custom_task_id': task,
        'custom_mode': 'interview', 'roles': 'Learner',
        'lis_outcome_service_url': 'http://example.com/outcomes', 'lis_result_sourcedid': task,
    }, timeout=30)
    assert result.status_code == 200
    assert foreign.get(host + '/api/interview/feedback/' + recording_id, timeout=20).status_code == 404
    assert foreign.get(host + '/api/interview/results/' + recording_id + '/', timeout=20).status_code == 404
